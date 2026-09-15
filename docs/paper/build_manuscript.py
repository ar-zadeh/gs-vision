"""Build and format the BRM tutorial manuscript.

Run from the repository root:

    python docs/paper/build_manuscript.py

Pandoc creates the Word document from ``paper.md``.  python-docx then applies
the APA-style manuscript layout and table formatting that Pandoc cannot express
through the reference document alone.
"""

from __future__ import annotations

import re
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PAPER_DIR = Path(__file__).resolve().parent
SOURCE = PAPER_DIR / "paper.md"
REFERENCE = PAPER_DIR / "reference.docx"
OUTPUT = PAPER_DIR / "gs-vision-BRM-tutorial.docx"


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_cell_margins(cell, points: float = 3.0) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    twips = str(int(points * 20))
    for edge in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), twips)
        node.set(qn("w:type"), "dxa")


def set_apa_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "insideH", "bottom"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "000000")
        borders.append(node)
    for edge in ("start", "end", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        borders.append(node)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, end))


def style_run(run, size: float = 12.0, font: str = "Times New Roman") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_paragraph(paragraph, *, double: bool = True, indent: bool = False) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 2.0 if double else 1.0
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.first_line_indent = Inches(0.5) if indent else Inches(0)
    for run in paragraph.runs:
        style_run(run)


def set_column_widths(table, widths: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(widths):
            row.cells[index].width = Inches(width)


def format_tables(document: Document) -> None:
    width_map = {
        2: [2.0, 4.5],
        3: [1.7, 2.8, 2.0],
        4: [1.9, 1.55, 1.55, 1.5],
        5: [2.6, 1.05, 1.05, 0.9, 0.9],
        7: [2.0, 0.9, 0.65, 0.7, 0.7, 0.8, 0.75],
    }
    for table in document.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_apa_borders(table)
        column_count = len(table.columns)
        if column_count in width_map:
            set_column_widths(table, width_map[column_count])
        font_size = 8.0 if column_count >= 7 or len(table.rows) > 15 else 9.0
        set_repeat_table_header(table.rows[0])
        for row_index, row in enumerate(table.rows):
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if row_index == 0
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.first_line_indent = Inches(0)
                    for run in paragraph.runs:
                        style_run(run, font_size)
                        if row_index == 0:
                            run.bold = True


def format_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.right_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.header_distance = Inches(0.5)
        header = section.header
        if not header.paragraphs:
            header.add_paragraph()
        paragraph = header.paragraphs[0]
        paragraph.clear()
        add_page_number(paragraph)
        for run in paragraph.runs:
            style_run(run)

    paragraph_styles = {
        "Normal",
        "Body Text",
        "First Paragraph",
        "Compact",
        "Bibliography",
    }
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name
        text = paragraph.text.strip()
        if style_name in paragraph_styles:
            indent = style_name in {"Body Text", "First Paragraph"}
            style_paragraph(paragraph, double=True, indent=indent)
        else:
            style_paragraph(paragraph, double=True, indent=False)

        if style_name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True
        elif style_name in {"Subtitle", "Author", "Date"}:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = False
        elif style_name == "Heading 1":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True
        elif style_name == "Heading 2":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_with_next = True
            for run in paragraph.runs:
                run.bold = True

        if text in {"Abstract", "Introduction", "References"} or text.startswith(
            "Appendix:"
        ):
            paragraph.paragraph_format.page_break_before = True

        if text.startswith("Keywords:"):
            paragraph.paragraph_format.first_line_indent = Inches(0)

        if re.fullmatch(r"Table \d+", text):
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = True
            if text == "Table 9":
                paragraph.paragraph_format.page_break_before = True
        elif paragraph.runs and all(
            run.italic or not run.text.strip() for run in paragraph.runs
        ):
            previous = paragraph._p.getprevious()
            previous_text = "" if previous is None else "".join(previous.itertext())
            if "Table " in previous_text:
                paragraph.paragraph_format.first_line_indent = Inches(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = True

        if text.startswith("Note."):
            paragraph.paragraph_format.first_line_indent = Inches(0)
            paragraph.paragraph_format.line_spacing = 1.0

        if style_name in {"Source Code", "Verbatim Char"}:
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.first_line_indent = Inches(0)
            for run in paragraph.runs:
                style_run(run, 8.0, "Consolas")

    in_abstract = False
    in_references = False
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "Abstract":
            in_abstract = True
            continue
        if in_abstract:
            paragraph.paragraph_format.first_line_indent = Inches(0)
            if text.startswith("Keywords:"):
                in_abstract = False
        if text == "References":
            in_references = True
            continue
        if in_references:
            if text.startswith("Appendix:"):
                in_references = False
            elif text:
                paragraph.paragraph_format.left_indent = Inches(0.5)
                paragraph.paragraph_format.first_line_indent = Inches(-0.5)

    format_tables(document)


def main() -> None:
    subprocess.run(
        [
            "pandoc",
            str(SOURCE),
            "--from=markdown+autolink_bare_uris",
            "--to=docx",
            f"--reference-doc={REFERENCE}",
            f"--resource-path={PAPER_DIR}",
            f"--output={OUTPUT}",
        ],
        check=True,
    )
    document = Document(OUTPUT)
    format_document(document)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
